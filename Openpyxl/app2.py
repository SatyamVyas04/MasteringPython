import openpyxl as xl
from openpyxl.chart import BarChart, Reference

# loading Workbook
wb = xl.load_workbook('Openpyxl/transactions2.xlsx')
# loading a singular sheet
sheet = wb["Sheet1"]

# adding a graph to the data
values = Reference(sheet, 
                   min_row=2, 
                   max_row=sheet.max_row,
                   min_col=4, 
                   max_col=4)

chart = BarChart()
chart.add_data(values)
sheet.add_chart(chart, "e2")

wb.save("transactions3.xlsx")